"""Source-first audit and partial replay; no claim to reproduce full author thermodynamics."""
from pathlib import Path
import json,hashlib,sys,zipfile
import numpy as np
import xarray as xr
from sklearn.mixture import GaussianMixture
from scipy.optimize import linear_sum_assignment
from openpyxl import load_workbook
p=Path(__file__).resolve().parent;data=p/'data';out=p/'research';out.mkdir(exist_ok=True)
for rid in ['20023151','21903512']:
 folder=data/rid
 if not folder.exists():
  with zipfile.ZipFile(data/(rid+'.zip')) as z:
   for name in z.namelist():assert (folder/name).resolve().is_relative_to(folder.resolve())
   z.extractall(folder)
src=data/'20023151';sys.path.insert(0,str(src));import single_shots_analysis_utils as aut
cases=[];audit=[]
for name in ['single_shots_three_cycles.nc','single_shots_repetitions.nc']:
 with xr.open_dataset(src/name,engine='h5netcdf') as d:
  re=d['real'].values;im=d['imag'].values;valid=np.isfinite(re)&np.isfinite(im);counts=valid.sum(axis=-1);a={'file':name,'dimensions':dict(d.sizes),'total_IQ_pairs':int(valid.size),'finite_IQ_pairs':int(valid.sum()),'empty_acquisition_cells':int(sum((counts==0).ravel())),'partial_acquisition_cells':int(sum(((counts>0)&(counts<re.shape[-1])).ravel())),'complete_acquisition_cells':int(sum((counts==re.shape[-1]).ravel())),'missing_fraction':float(1-valid.mean())};audit.append(a)
  cases.append({'eye':'data_audit','kind':'acquisition','source':'Zenodo 20023151 / '+name,'case':{'total_entries':a['total_IQ_pairs'],'finite_entries':a['finite_IQ_pairs'],'stage':'raw IQ extracts with unacquired coordinate combinations padded','calibration_supplied':True,'energy_schedule_supplied':False,'state_assignment_verified':False}})
(out/'raw_audit.json').write_text(json.dumps(audit,indent=2));print('raw audit',audit,flush=True)
# Authors' prepared-state selections; deterministic label assignment replaces random retry loop.
with xr.open_dataset(src/'single_shots_calibration.nc',engine='h5netcdf') as d:
 X=[]
 for st,ef in [(0,0),(1,0),(2,1)]:
  z=d.isel(state=st,pi_ef_length=ef);x=np.column_stack([z['real'].values.ravel(),z['imag'].values.ravel()]);assert np.isfinite(x).all();X.append(x)
fit=GaussianMixture(n_components=4,covariance_type='full',random_state=20260907,n_init=5).fit(np.concatenate(X))
means=np.array([x.mean(axis=0) for x in X]);_,idx=linear_sum_assignment(np.linalg.norm(means[:,None]-fit.means_[None,:],axis=-1));order=list(idx)+[j for j in range(4) if j not in idx];gm=fit.means_[order];gc=fit.covariances_[order]
# Original ellipse helper; ten million samples per prepared Gaussian, chunked to cap memory.
rng=np.random.default_rng(20260907);corr=np.zeros((4,4));narrow=np.zeros((4,4));N=10000000
for st in range(4):
 for chunk in range(N//100000):
  x=rng.multivariate_normal(gm[st],gc[st],100000)
  for region in range(4):
   corr[st,region]+=len(aut.points_in_ellipse(x[:,0],x[:,1],gm[region],gc[region],1)[0])
   narrow[st,region]+=len(aut.points_in_ellipse(x[:,0],x[:,1],gm[region],.4*gc[region],1)[0])
 print('calibration sampled state',st,flush=True)
corr/=N;narrow/=N
(out/'readout_calibration.json').write_text(json.dumps({'means':gm.tolist(),'covariances':gc.tolist(),'original_scale_selection_matrix':corr.tolist(),'matched_narrow_selection_matrix':narrow.tolist(),'samples_per_gaussian':N,'changes_from_notebook':['Seeded five-start GMM and deterministic prepared-mean label assignment replace random index retry.','Chunked Monte Carlo with same 10 million samples/state.','No plotting or full Qutip thermodynamic simulation.'],'original_map_convention':'corr[prepared_state, readout_region]','notebook_application':'counts = corr @ counts, then normalize; acquisition covariance scale 0.4 while original corr sampling uses scale 1.0','inference_status':'Partial pipeline replay plus forward-model diagnostic; no corrected efficiency asserted.'},indent=2))
records=[]
for name in ['single_shots_three_cycles.nc','single_shots_repetitions.nc']:
 with xr.open_dataset(src/name,engine='h5netcdf') as d:
  re=d['real'].values;im=d['imag'].values
  for ix in np.ndindex(re.shape[:-1]):
   ok=np.isfinite(re[ix])&np.isfinite(im[ix])
   if not ok.any():continue
   x=np.column_stack([re[ix][ok],im[ix][ok]]);counts=np.array([len(aut.points_in_ellipse(x[:,0],x[:,1],gm[j],.4*gc[j],1)[0]) for j in range(4)]);replayed=corr@counts;inverse=np.linalg.solve(narrow.T,counts)
   records.append({'file':name,'index':list(ix),'coordinates':{dim:float(d.coords[dim].values[j]) for dim,j in zip(d['real'].dims[:-1],ix)},'shots':len(x),'selected_counts':counts.tolist(),'notebook_replay_populations':(replayed/sum(replayed)).tolist(),'matched_inverse_populations':(inverse/sum(inverse)).tolist()})
(out/'raw_population_replay.json').write_text(json.dumps(records,indent=2));cases.append({'eye':'readout_transfer','kind':'processed','source':'Seeded partial replay of Zenodo 20023151 calibration and first acquired cell','case':{'selection_matrix':narrow.tolist(),'selected_counts':records[0]['selected_counts']}})
# Read-only author spreadsheets. Preserve Data vs Model labels, coordinates, and errors.
root=data/'21903512'/'data'/'1_Main_Text'
w=load_workbook(root/'Figure_3/data_fig_3.xlsx',data_only=True,read_only=True);rows=list(w.active.values);diags=np.array([[float(v) for v in row] for row in [[rr[1],rr[2],rr[3],rr[7],rr[8],rr[9]] for rr in rows[3:10]]]);models=[]
for side in ['heating','cooling']:
 for col,T in enumerate([.5,1.5,2.5]):
  dd=diags[:,col+(0 if side=='heating' else 3)];K=np.diag(dd)
  for j in range(7):
   if dd[j]<0:K[j+(-1 if side=='heating' else 1),j]=-dd[j]
  P=[0]*7;P[6 if side=='heating' else 0]=1;c={'energies':list(range(6,-1,-1)),'generator':K.tolist(),'initial':P,'times':np.linspace(0,3,61).tolist()};cases.append({'eye':'ladder','kind':'theoretical','source':'Burgardt et al. figure 3 published MODEL rates, T='+str(T)+' uK, '+side,'case':c});models.append({'temperature_uK':T,'stroke':side,'rate_eigenvalues_s_inv':dd.tolist()})
w=load_workbook(root/'Figure_7/data_fig_7.xlsx',data_only=True,read_only=True);rows=list(w.active.values);points=[]
for rr in rows[3:]:
 if isinstance(rr[0],(int,float)) and isinstance(rr[13],(int,float)):
  points.append({'temperature_uK':rr[0],'cycle_duration_s':rr[1],'work_numeric':rr[7],'work_error_numeric':rr[8],'power_kB_nK_per_ms':rr[14],'power_error':rr[15],'W_over_cycle_minus_power':rr[7]/rr[1]-rr[14]})
(out/'atomic_source_results.json').write_text(json.dumps({'source':'https://zenodo.org/records/21903512','stage':'published figure tables; NOT raw detector events','figure3':'Model eigenvalues, directed nearest-neighbor generator reconstructed using paper equations 7–8; energy units arbitrary level gap in eye.','models':models,'figure7_experimental_points':points,'unit_issue':'Figure 7 work column labels nK/ms although it is a work quantity; numeric W/t matches power. Work units need confirmation from author before SI conversion.','derived_peak_ratio_1140_over_679':points[2]['power_kB_nK_per_ms']/points[0]['power_kB_nK_per_ms'] if len(points)>=3 else None},indent=2))
(out/'source_eye_cases.json').write_text(json.dumps(cases,indent=2));print('finished',len(records),'raw acquired cells;',len(cases),'source eye cases',flush=True)
